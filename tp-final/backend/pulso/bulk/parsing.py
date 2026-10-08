"""Turn an uploaded .csv/.txt/.xlsx into rows keyed by normalised header. No database access."""

import csv
import io
import zipfile
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from ..errors import APIError
from .cells import blank_to_none, normalize_header

DELIMITERS = ',;\t|'
MAX_BLANK_RUN = 10_000  # consecutive blank lines after which scanning stops (styled-but-empty tails)
MAX_COLUMNS = 100  # widest header row accepted
MAX_UNCOMPRESSED = 50 * 1024 * 1024  # bytes an .xlsx may expand to


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[tuple[int, dict[str, Any]]]  # (line number in the file, {header: raw cell})


def invalid(message: str, code: str = 'archivo_invalido') -> APIError:
    return APIError(message, 400, code)


def text_rows(content: bytes) -> Iterator[tuple[int, list[Any]]]:
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('cp1252', errors='replace')
    if '\x00' in text:
        raise invalid('El archivo contiene caracteres nulos.')
    first = next((line for line in text.splitlines() if line.strip()), '')
    delimiter = max(DELIMITERS, key=first.count) if any(d in first for d in DELIMITERS) else ','
    reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True)
    try:
        for cells in reader:
            yield reader.line_num, cells
    except csv.Error:
        raise invalid('El archivo de texto no se pudo leer: revisá comillas y delimitadores.') from None


def xlsx_rows(content: bytes) -> Iterator[tuple[int, list[Any]]]:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise invalid('El archivo .xlsx no es válido.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
    except Exception:  # corrupt central directory
        raise invalid('El archivo .xlsx no es válido.') from None
    if expanded > MAX_UNCOMPRESSED:
        raise invalid('El archivo .xlsx es demasiado grande una vez descomprimido.')
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        # A declared <dimension> of A1:XFD1048576 would make openpyxl pad every missing row to 16384 columns.
        sheet.reset_dimensions()
    except Exception:  # openpyxl raises many types for broken workbooks
        raise invalid('El archivo .xlsx no es válido.') from None
    try:
        # max_col bounds the padding openpyxl adds up to a row's last cell (one extra column so that a header
        # wider than MAX_COLUMNS is still detected as too wide).
        rows = iter(sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1))
        number = 0
        while True:
            try:
                row = next(rows)
            except StopIteration:
                break
            except Exception:  # truncated/corrupt sheet XML surfaces while iterating
                raise invalid('El archivo .xlsx no es válido.') from None
            number += 1
            # C-level shortcut for rows with only empty cells; whitespace-only cells are handled downstream.
            yield number, [] if row.count(None) == len(row) else list(row)
    finally:
        workbook.close()


def parse_file(filename: str, content: bytes, max_rows: int) -> ParsedFile:
    name = filename.lower()
    if name.endswith('.xlsx'):
        lines = xlsx_rows(content)
    elif name.endswith(('.csv', '.txt')):
        lines = text_rows(content)
    else:
        raise APIError(
            'Formato no soportado. Subí un archivo .csv, .xlsx o .txt.', 400, 'formato_no_soportado'
        )

    headers: list[str] | None = None
    rows: list[tuple[int, dict[str, Any]]] = []
    blank_run = 0
    for line, cells in lines:
        if all(blank_to_none(cell) is None for cell in cells):
            blank_run += 1
            if blank_run > MAX_BLANK_RUN:
                break
            continue
        blank_run = 0
        if headers is None:
            end = len(cells)
            while end and blank_to_none(cells[end - 1]) is None:
                end -= 1
            cells = cells[:end]
            if len(cells) > MAX_COLUMNS:
                raise invalid(f'El archivo tiene más de {MAX_COLUMNS} columnas.')
            headers = ['' if cell is None else normalize_header(cell) for cell in cells]
            counts = Counter(h for h in headers if h)
            duplicated = sorted(h for h, count in counts.items() if count > 1)
            if duplicated:
                raise invalid('Hay columnas repetidas: ' + ', '.join(duplicated) + '.')
            continue
        if len(rows) >= max_rows:
            raise invalid(
                f'El archivo supera el máximo de {max_rows} filas para esta carga.',
                'archivo_demasiado_grande',
            )
        rows.append(
            (line, {headers[i]: cells[i] for i in range(min(len(headers), len(cells))) if headers[i]})
        )
    if headers is None:
        raise invalid('El archivo está vacío.')
    if not rows:
        raise invalid('El archivo no tiene filas de datos debajo de los encabezados.')
    return ParsedFile([h for h in headers if h], rows)
